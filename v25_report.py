"""V2.5 report wrapper: preserves V2 workbook and adds audit/history/backtest tabs."""
from pathlib import Path
from datetime import date
import csv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from excel_report_v2 import build_excel_v2
from backtest import evaluate

def _read_history(path="data/history.csv"):
    p = Path(path)
    if not p.exists():
        return []
    with p.open("r", encoding="utf-8") as f:
        return list(csv.DictReader(f))

def build_excel_v25(rows, all_cf, output_path, history_path="data/history.csv"):
    build_excel_v2(rows, all_cf, output_path)
    wb = load_workbook(output_path)

    # Rename the main title to make the version explicit.
    ws = wb["Executive Summary"]
    ws["A1"] = "EGX100 V2.5 Fundamental Research Report"
    ws["A3"] = ws["A3"].value.replace("V2 sector-aware", "V2.5 sector-aware") if ws["A3"].value else "V2.5"

    audit = wb.create_sheet("Data Quality")
    audit.append(["Ticker","Sector","Data Completeness","Methods","Confidence","Confidence Score","Valuation Dispersion","Scraper Errors"])
    for cell in audit[1]:
        cell.font = Font(bold=True,color="FFFFFF")
        cell.fill = PatternFill("solid",fgColor="1F4E78")
    for r in rows:
        audit.append([
            r["ticker"],r["sector"],r["data_completeness"]/100,r["methods_used"],
            r["confidence"],r["confidence_score"],r["dispersion"],
            "; ".join(r["errors"])
        ])
    for row in audit.iter_rows(min_row=2,min_col=3,max_col=3):
        row[0].number_format="0%"

    hist = wb.create_sheet("Historical Signals")
    hrows = _read_history(history_path)
    if hrows:
        hist.append(list(hrows[0].keys()))
        for r in hrows:
            hist.append(list(r.values()))
    else:
        hist.append(["No historical snapshots yet. The first successful run creates data/history.csv."])

    bt = wb.create_sheet("Backtest")
    summary, observations = evaluate(history_path)
    bt.append(["Horizon","Observations","Scored Observations","Average Return","Median Return","Positive Rate","Average Opportunity Score"])
    for s in summary:
        bt.append([s["horizon"],s["observations"],s["scored_observations"],s["avg_return"],s["median_return"],s["positive_rate"],s["avg_score"]])
    if not summary:
        bt.append(["Not enough future observations yet. Backtest activates automatically as weekly history accumulates."])
    for row in bt.iter_rows(min_row=2,min_col=4,max_col=6):
        for cell in row: cell.number_format="0.0%"

    ass = wb.create_sheet("Assumptions & Audit")
    ass.column_dimensions["A"].width=32
    ass.column_dimensions["B"].width=100
    items = [
        ("Report version","V2.5"),
        ("Generated",date.today().isoformat()),
        ("Sector classification","Manual overrides -> controlled keyword taxonomy -> Unclassified; unknown names are not forced into a sector."),
        ("Financial valuation","P/E + peer P/B; corporate DCF/EV-EBITDA are excluded."),
        ("Real-estate valuation","P/E/P/B + DCF + EV/EBITDA where valid; full NAV is a V3 target."),
        ("Corporate valuation","P/E + DCF + EV/EBITDA where valid."),
        ("Fair value","Median of valid methods; bear/base/bull range; method dispersion is explicitly reported."),
        ("Opportunity Score","Current V2 weights remain research assumptions and are not treated as statistically validated until backtesting."),
        ("History","Each successful weekly run stores the signal snapshot by ticker/date in data/history.csv."),
        ("Backtest","1-week, 4-week and 12-week realized returns are measured only when future snapshots exist."),
        ("Data provenance","Current fundamental source remains StockAnalysis scraping; official-company/EGX reconciliation is required before client distribution."),
        ("Research status","Quantitative research/screening only; not investment advice.")
    ]
    for i,(a,b) in enumerate(items,1):
        ass.cell(i,1,a).font=Font(bold=True)
        ass.cell(i,2,b)
        ass.cell(i,2).alignment=Alignment(wrap_text=True,vertical="top")

    wb.save(output_path)
    return output_path
