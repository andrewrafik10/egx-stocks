"""
Builds the weekly Excel deliverable with a more professional structure:
  1. Executive Summary
  2. Full Ranking
  3. Methodology
"""

from datetime import date
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FINANCIAL_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
HIGH_CONF_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
MEDIUM_CONF_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
LOW_CONF_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=12)

COLUMNS = [
    ("Rank", 6),
    ("Ticker", 10),
    ("Type", 10),
    ("Sector", 20),
    ("Price", 10),
    ("EPS", 10),
    ("BVPS", 10),
    ("Revenue", 14),
    ("Net Income", 14),
    ("EBITDA", 14),
    ("P/E Fair Value", 14),
    ("P/E Upside", 12),
    ("Graham Fair Value", 16),
    ("Graham Upside", 14),
    ("DCF Fair Value", 14),
    ("DCF Upside", 12),
    ("Comps Fair Value", 14),
    ("Comps Upside", 12),
    ("Composite Upside", 16),
    ("Methods Used", 12),
    ("Confidence", 12),
    ("Notes", 45),
]

PERCENT_COLS = {"P/E Upside", "Graham Upside", "DCF Upside", "Comps Upside", "Composite Upside"}
NUMBER_COLS = {"Price", "EPS", "BVPS", "P/E Fair Value", "Graham Fair Value", "DCF Fair Value", "Comps Fair Value"}
BIG_NUMBER_COLS = {"Revenue", "Net Income", "EBITDA"}


def build_excel(ranked: list, all_cf: dict, output_path: str) -> str:
    wb = Workbook()

    # ====================== 1. Executive Summary ======================
    ws_sum = wb.active
    ws_sum.title = "Executive Summary"

    usable = [r for r in ranked if r["composite_upside"] is not None]
    high_conf = [r for r in usable if r.get("confidence") == "High"]
    medium_conf = [r for r in usable if r.get("confidence") == "Medium"]
    low_conf = [r for r in usable if r.get("confidence") == "Low"]

    ws_sum["A1"] = "EGX Weekly Fundamental Report"
    ws_sum["A1"].font = Font(bold=True, size=16)
    ws_sum["A2"] = f"Generated: {date.today().strftime('%d %B %Y')}"
    ws_sum["A3"] = f"Universe: {len(ranked)} tickers | Scored: {len(usable)}"

    ws_sum["A5"] = "Confidence Breakdown"
    ws_sum["A5"].font = SECTION_FONT
    ws_sum["A6"] = f"High Confidence (≥3 methods): {len(high_conf)}"
    ws_sum["A7"] = f"Medium Confidence (2 methods): {len(medium_conf)}"
    ws_sum["A8"] = f"Low Confidence (≤1 method): {len(low_conf)}"

    ws_sum["A10"] = "Top 10 Opportunities (Highest Composite Upside)"
    ws_sum["A10"].font = SECTION_FONT

    headers = ["Rank", "Ticker", "Sector", "Price", "Composite Upside", "Methods", "Confidence"]
    for col, h in enumerate(headers, 1):
        cell = ws_sum.cell(row=11, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for i, r in enumerate(usable[:10], 1):
        row = 11 + i
        ws_sum.cell(row=row, column=1, value=i)
        ws_sum.cell(row=row, column=2, value=r["ticker"])
        ws_sum.cell(row=row, column=3, value=r.get("macro_sector", ""))
        ws_sum.cell(row=row, column=4, value=r["price"]).number_format = "#,##0.00"
        cell_up = ws_sum.cell(row=row, column=5, value=r["composite_upside"])
        cell_up.number_format = "+0.0%;-0.0%"
        ws_sum.cell(row=row, column=6, value=r["methods_used"])
        conf_cell = ws_sum.cell(row=row, column=7, value=r.get("confidence", ""))
        if r.get("confidence") == "High":
            conf_cell.fill = HIGH_CONF_FILL
        elif r.get("confidence") == "Medium":
            conf_cell.fill = MEDIUM_CONF_FILL
        else:
            conf_cell.fill = LOW_CONF_FILL

    # Sector average upside
    ws_sum["A23"] = "Average Composite Upside by Sector"
    ws_sum["A23"].font = SECTION_FONT

    sector_upsides = defaultdict(list)
    for r in usable:
        sector_upsides[r.get("macro_sector", "Other")].append(r["composite_upside"])

    ws_sum["A24"] = "Sector"
    ws_sum["B24"] = "Avg Upside"
    ws_sum["C24"] = "Count"
    ws_sum["A24"].fill = HEADER_FILL
    ws_sum["B24"].fill = HEADER_FILL
    ws_sum["C24"].fill = HEADER_FILL
    ws_sum["A24"].font = HEADER_FONT
    ws_sum["B24"].font = HEADER_FONT
    ws_sum["C24"].font = HEADER_FONT

    row = 25
    for sector, ups in sorted(sector_upsides.items(), key=lambda x: -sum(x[1])/len(x[1])):
        avg = sum(ups) / len(ups)
        ws_sum.cell(row=row, column=1, value=sector)
        cell = ws_sum.cell(row=row, column=2, value=avg)
        cell.number_format = "+0.0%;-0.0%"
        ws_sum.cell(row=row, column=3, value=len(ups))
        row += 1

    ws_sum.column_dimensions["A"].width = 28
    ws_sum.column_dimensions["B"].width = 14
    ws_sum.column_dimensions["C"].width = 12
    ws_sum.column_dimensions["D"].width = 12
    ws_sum.column_dimensions["E"].width = 16
    ws_sum.column_dimensions["F"].width = 10
    ws_sum.column_dimensions["G"].width = 12

    # ====================== 2. Full Ranking ======================
    ws = wb.create_sheet("Full Ranking")

    for col_idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "C2"

    col_index = {name: i + 1 for i, (name, _) in enumerate(COLUMNS)}

    for rank, r in enumerate(ranked, start=1):
        row = rank + 1
        cf = all_cf.get(r["ticker"])
        mr = r["method_results"]

        def fv(method):
            return mr.get(method, {}).get("fair_value")

        def up(method):
            return r["upsides"].get(method)

        note_parts = list(r.get("errors", []))
        if r.get("low_confidence"):
            note_parts.insert(0, f"LOW CONFIDENCE ({r['methods_used']} method{'s' if r['methods_used'] != 1 else ''} only)")

        values = {
            "Rank": rank,
            "Ticker": r["ticker"],
            "Type": "Financial" if r["is_financial"] else "Other",
            "Sector": r.get("macro_sector") or (r.get("sector") or ""),
            "Price": r["price"],
            "EPS": cf.eps if cf else None,
            "BVPS": cf.book_value_per_share if cf else None,
            "Revenue": cf.revenue if cf else None,
            "Net Income": cf.net_income if cf else None,
            "EBITDA": cf.ebitda if cf else None,
            "P/E Fair Value": fv("pe"),
            "P/E Upside": up("pe"),
            "Graham Fair Value": fv("graham"),
            "Graham Upside": up("graham"),
            "DCF Fair Value": fv("dcf"),
            "DCF Upside": up("dcf"),
            "Comps Fair Value": fv("comps"),
            "Comps Upside": up("comps"),
            "Composite Upside": r["composite_upside"],
            "Methods Used": r["methods_used"],
            "Confidence": r.get("confidence", ""),
            "Notes": "; ".join(note_parts) if note_parts else "",
        }

        for name, value in values.items():
            c = ws.cell(row=row, column=col_index[name], value=value)
            c.border = THIN_BORDER

            if name in PERCENT_COLS and value is not None:
                c.number_format = "+0.0%;-0.0%"
            elif name in NUMBER_COLS and value is not None:
                c.number_format = "#,##0.00"
            elif name in BIG_NUMBER_COLS and value is not None:
                c.number_format = '#,##0,,"M"'

            if r["is_financial"]:
                c.fill = FINANCIAL_FILL

            if name == "Confidence":
                if value == "High":
                    c.fill = HIGH_CONF_FILL
                elif value == "Medium":
                    c.fill = MEDIUM_CONF_FILL
                elif value == "Low":
                    c.fill = LOW_CONF_FILL

    last_row = len(ranked) + 1
    if last_row > 1:
        for name in PERCENT_COLS:
            col_letter = get_column_letter(col_index[name])
            rng = f"{col_letter}2:{col_letter}{last_row}"
            ws.conditional_formatting.add(
                rng,
                ColorScaleRule(
                    start_type="min", start_color="F8696B",
                    mid_type="percentile", mid_value=50, mid_color="FFEB84",
                    end_type="max", end_color="63BE7B",
                ),
            )

    # ====================== 3. Methodology ======================
    notes = wb.create_sheet("Methodology")
    notes["A1"] = "EGX Weekly Fundamental Report – Methodology"
    notes["A1"].font = Font(bold=True, size=14)

    notes["A3"] = "Data scope: EGX100 universe. Values are based on the latest successfully scraped fundamentals available at run time."
    notes["A4"] = "Valuation methods: P/E, Graham Number, DCF, and EV/EBITDA comparables where applicable."
    notes["A5"] = "Composite upside is calculated from the valuation methods that have valid inputs."
    notes["A6"] = "Confidence reflects the number of valuation methods successfully available for a stock."
    notes["A8"] = "Important: DCF cost of equity currently uses configured sector defaults and should be upgraded to documented CAPM-based assumptions before commercial use."
    notes["A9"] = "Data exceptions and missing fields are shown in the Full Ranking Notes column."
    notes.column_dimensions["A"].width = 110
    notes.column_dimensions["B"].width = 20
    notes.column_dimensions["C"].width = 20

    # Save the workbook to the exact path supplied by main.py.
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
