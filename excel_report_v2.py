"""Commercial-style V2 Excel output."""
from datetime import date
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

NAVY="1F4E78"; WHITE="FFFFFF"; GREEN="C6EFCE"; YELLOW="FFEB9C"; RED="FFC7CE"
HEADER=PatternFill("solid",fgColor=NAVY)
HF=Font(color=WHITE,bold=True)
BORDER=Border(*(Side(style="thin",color="D9D9D9"),)*4)

def _header(ws,row,headers):
    for i,h in enumerate(headers,1):
        c=ws.cell(row,i,h); c.fill=HEADER; c.font=HF; c.alignment=Alignment(horizontal="center")
        c.border=BORDER

def _pct(v):
    return None if v is None else v

def build_excel_v2(rows, all_cf, output_path):
    wb=Workbook()
    ws=wb.active; ws.title="Executive Summary"
    scored=[r for r in rows if r["opportunity_score"] is not None]
    ws["A1"]="EGX100 V2 Fundamental Research Report"; ws["A1"].font=Font(bold=True,size=18)
    ws["A2"]=f"Generated: {date.today():%d %B %Y}"
    ws["A3"]=f"Universe: {len(rows)} | Scored: {len(scored)} | V2 sector-aware valuation + quality + dispersion"
    ws["A5"]="V2 Market Snapshot"; ws["A5"].font=Font(bold=True,size=13)
    stats=[
        ("Stocks scored",len(scored)),
        ("High confidence",sum(r["confidence"]=="High" for r in scored)),
        ("Medium confidence",sum(r["confidence"]=="Medium" for r in scored)),
        ("Low confidence",sum(r["confidence"]=="Low" for r in scored)),
        ("Median base upside", sorted([(r["fair_value_base"]-r["price"])/r["price"] for r in scored if r["fair_value_base"] and r["price"]])[len([r for r in scored if r["fair_value_base"] and r["price"]])//2] if any(r["fair_value_base"] and r["price"] for r in scored) else None),
    ]
    for i,(k,v) in enumerate(stats,7):
        ws.cell(i,1,k); ws.cell(i,2,v)
        if "upside" in k.lower() and v is not None: ws.cell(i,2).number_format="+0.0%;-0.0%"
    ws["A14"]="Top 10 V2 Opportunities"; ws["A14"].font=Font(bold=True,size=13)
    headers=["Rank","Ticker","Sector","Price","Base FV","Base Upside","Bear FV","Bull FV","Dispersion","Quality","Opportunity","Confidence","Data Completeness","Methods"]
    _header(ws,15,headers)
    for i,r in enumerate(scored[:10],16):
        vals=[r["rank"],r["ticker"],r["sector"],r["price"],r["fair_value_base"],
              (r["fair_value_base"]-r["price"])/r["price"] if r["fair_value_base"] and r["price"] else None,
              r["fair_value_bear"],r["fair_value_bull"],r["dispersion"],r["quality_score"],
              r["opportunity_score"],r["confidence"],r["data_completeness"]/100,r["methods_used"]]
        for j,v in enumerate(vals,1): ws.cell(i,j,v).border=BORDER
        for j in (6,): ws.cell(i,j).number_format="+0.0%;-0.0%"
        for j in (9,): ws.cell(i,j).number_format="0.0%"
        ws.cell(i,13).number_format="0%"
    for col,w in enumerate([7,10,22,11,12,13,12,12,12,10,12,13,18,10],1):
        ws.column_dimensions[get_column_letter(col)].width=w
    ws.freeze_panes="A16"

    # Sector dashboard
    sec=wb.create_sheet("Sector Dashboard"); sec["A1"]="Sector Dashboard"; sec["A1"].font=Font(bold=True,size=16)
    sh={}
    for r in scored: sh.setdefault(r["sector"],[]).append(r)
    _header(sec,3,["Sector","Stocks","Median Base Upside","Avg Opportunity","Avg Quality","Avg Confidence","High Confidence"])
    for i,(s,items) in enumerate(sorted(sh.items()),4):
        ups=sorted([(r["fair_value_base"]-r["price"])/r["price"] for r in items if r["fair_value_base"] and r["price"]])
        sec.cell(i,1,s); sec.cell(i,2,len(items))
        if ups: sec.cell(i,3,ups[len(ups)//2]); sec.cell(i,3).number_format="+0.0%"
        sec.cell(i,4,sum(r["opportunity_score"] for r in items)/len(items))
        q=[r["quality_score"] for r in items if r["quality_score"] is not None]
        sec.cell(i,5,sum(q)/len(q) if q else None)
        sec.cell(i,6,sum(r["confidence_score"] for r in items)/len(items))
        sec.cell(i,7,sum(r["confidence"]=="High" for r in items))
    chart=BarChart(); chart.title="Median Base Upside by Sector"; chart.y_axis.title="Upside"; chart.x_axis.title="Sector"
    chart.add_data(Reference(sec,min_col=3,min_row=3,max_row=3+len(sh)),titles_from_data=True)
    chart.set_categories(Reference(sec,min_col=1,min_row=4,max_row=3+len(sh))); sec.add_chart(chart,"I3")

    # Full ranking
    fr=wb.create_sheet("V2 Ranking")
    headers=["Rank","Ticker","Sector","Price","Bear FV","Base FV","Bull FV","Base Upside","Valuation Dispersion","Quality Score","Opportunity Score","Confidence","Confidence Score","Data Completeness","Methods","P/E","P/B","DCF","EV/EBITDA","ROE","Net Margin","EBITDA Margin","FCF Margin","Debt/Equity","EPS CAGR","Revenue CAGR","Peer P/E","Peer P/B","Peer EV/EBITDA","Data Notes"]
    _header(fr,1,headers); fr.freeze_panes="D2"
    for i,r in enumerate(rows,2):
        m=r["metrics"]; p=r["peer"]; vals=[
            r["rank"],r["ticker"],r["sector"],r["price"],r["fair_value_bear"],r["fair_value_base"],r["fair_value_bull"],
            (r["fair_value_base"]-r["price"])/r["price"] if r["fair_value_base"] and r["price"] else None,
            r["dispersion"],r["quality_score"],r["opportunity_score"],r["confidence"],r["confidence_score"],r["data_completeness"]/100,r["methods_used"],
            m["pe"],m["pb"],r["valuations"].get("dcf"),r["valuations"].get("comps"),m["roe"],m["net_margin"],m["ebitda_margin"],m["fcf_margin"],m["debt_equity"],m["eps_growth"],m["revenue_growth"],p.get("pe"),p.get("pb"),p.get("ev_ebitda"),"; ".join(r["errors"])
        ]
        for j,v in enumerate(vals,1): fr.cell(i,j,v).border=BORDER
        for j in (8,9,20,21,22,23,25,26): fr.cell(i,j).number_format="0.0%"
        fr.cell(i,14).number_format="0%"
        conf=fr.cell(i,12)
        conf.fill=PatternFill("solid",fgColor=GREEN if r["confidence"]=="High" else YELLOW if r["confidence"]=="Medium" else RED)
    widths=[7,10,20,10,11,11,11,13,16,12,14,12,15,16,9,9,9,10,11,10,12,14,12,12,11,11,11,11,14,35]
    for j,w in enumerate(widths,1): fr.column_dimensions[get_column_letter(j)].width=w
    fr.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
    for col in (8,9,11,20,21,22,23,25,26):
        fr.conditional_formatting.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{len(rows)+1}",ColorScaleRule(start_type="min",start_color="F8696B",mid_type="percentile",mid_value=50,mid_color="FFEB84",end_type="max",end_color="63BE7B"))

    # Valuation detail
    vd=wb.create_sheet("Valuation Detail")
    _header(vd,1,["Ticker","Sector","Price","P/E FV","P/B FV","DCF FV","EV/EBITDA FV","Bear FV","Base FV","Bull FV","Base Upside","Dispersion","Methods","Peer P/E","Peer P/B","Peer EV/EBITDA"])
    for i,r in enumerate(rows,2):
        v=r["valuations"]; p=r["peer"]
        vals=[r["ticker"],r["sector"],r["price"],v.get("pe"),v.get("pb"),v.get("dcf"),v.get("comps"),r["fair_value_bear"],r["fair_value_base"],r["fair_value_bull"],(r["fair_value_base"]-r["price"])/r["price"] if r["fair_value_base"] and r["price"] else None,r["dispersion"],r["methods_used"],p.get("pe"),p.get("pb"),p.get("ev_ebitda")]
        for j,x in enumerate(vals,1): vd.cell(i,j,x).border=BORDER
        for j in (11,12): vd.cell(i,j).number_format="0.0%"

    # Methodology
    mt=wb.create_sheet("Methodology"); mt.column_dimensions["A"].width=115
    mt["A1"]="V2 Methodology"; mt["A1"].font=Font(bold=True,size=16)
    notes=[
        "V2 is an automated quantitative research/screening model, not investment advice.",
        "Financials: P/E and peer P/B are emphasized; generic corporate DCF/EV-EBITDA are excluded from the primary blend.",
        "Real Estate: P/E/PB, DCF and EV/EBITDA may be used when inputs are valid. P/B is not a substitute for a true property NAV model.",
        "Other sectors: P/E, DCF and EV/EBITDA are used where valid.",
        "Base fair value = median of valid method outputs. Bear/Bull use robust quartiles when 3+ methods exist; otherwise a 15% range around the sole method.",
        "Valuation dispersion = (max method FV - min method FV) / base FV. High dispersion reduces confidence.",
        "Quality score uses available ROE, net margin, FCF conversion, leverage and EPS growth. Missing data lowers confidence.",
        "Opportunity Score combines valuation (50%), quality (25%) and confidence (25%). These weights are model assumptions and should be backtested before commercial use.",
        "Peer multiples are calculated from the current EGX100 scrape within broad macro sectors. Extreme peer multiples are excluded from medians.",
        "DCF currently uses the configured cost-of-equity fallback. Production V3 should use stock-level CAPM inputs with documented beta/risk-free/ERP timestamps.",
        "Sector classification uses scraped industry text plus controlled keyword mapping. Unclassified names are not forced into a sector.",
        "All source data should be timestamped and reconciled with official EGX/company filings before client distribution."
    ]
    for i,n in enumerate(notes,3): mt.cell(i,1,"• "+n)
    output=Path(output_path); output.parent.mkdir(parents=True,exist_ok=True); wb.save(output_path); return output_path
